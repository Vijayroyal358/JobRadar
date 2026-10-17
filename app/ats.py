"""ATS (applicant tracking system) checks and the ATS-safe resume builder.

Two jobs: score a tailored resume against the rules ATS parsers care about (single column, no
tables or text boxes, contact details in the body, standard headings, safe fonts, job keywords
present) and build a plain single-column DOCX that any parser reads correctly.
"""
import re
from pathlib import Path

from .resume_parser import EXP_HEAD, OTHER_HEAD, RANGE

SAFE_FONTS = {"arial", "calibri", "cambria", "georgia", "garamond", "helvetica", "times new roman", "verdana",
              "tahoma", "trebuchet ms", "book antiqua", "palatino linotype", "century gothic", "lato", "roboto",
              "open sans", "segoe ui", "aptos", "liberation sans", "liberation serif", "dejavu sans", "carlito"}
EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE = re.compile(r"(?:\+?\d[\d\s\-().]{7,}\d)")
LINK = re.compile(r"(?:https?://|www\.|linkedin\.com/|github\.com/)\S+", re.I)
SAFE_BULLETS = "•▪◦●‣■–—’‘“”·⁃ →…"
ODD_CHARS = re.compile(r"[^\x09\x0a\x0d\x20-\x7e¡-ɏ" + SAFE_BULLETS + r"]")
BULLET = re.compile(r"^\s*([•▪◦●‣■⁃\-\*·–])\s+")

# Standard section names ATS parsers map to their internal fields, keyed by the text we see.
SECTIONS = {
    "summary": (re.compile(r"^(professional\s+)?(summary|profile|objective|about(\s+me)?|career\s+(summary|objective|profile))\s*:?$", re.I), "Professional Summary"),
    "experience": (EXP_HEAD, "Work Experience"),
    "education": (re.compile(r"^(education(al)?(\s+(background|qualifications?|details))?|academics?(\s+\w+)?|qualifications?)\s*:?$", re.I), "Education"),
    "skills": (re.compile(r"^((technical|key|core|professional)\s+)?(skills|competencies|technologies|tech\s+stack|expertise)(\s+(and|&)\s+\w+)?\s*:?$", re.I), "Skills"),
    "certifications": (re.compile(r"^(certifications?|licenses?(\s+(and|&)\s+certifications?)?|trainings?(\s+(and|&)\s+certifications?)?)\s*:?$", re.I), "Certifications"),
    "projects": (re.compile(r"^((academic|personal|key)\s+)?projects?\s*:?$", re.I), "Projects"),
}
REQUIRED_SECTIONS = ("experience", "education", "skills")


def section_of(text):
    t = (text or "").strip()
    if not t or len(t) > 60:
        return None
    for key, (rx, _) in SECTIONS.items():
        if rx.match(t):
            return key
    if OTHER_HEAD.match(t):
        return "other"
    return None


# ---------------------------------------------------------------- keyword coverage
def _kw_regex(kw):
    k = re.escape(kw.strip())
    k = k.replace(r"\ ", r"[\s\-/]*")
    return re.compile(r"(?<![A-Za-z0-9+#.])" + k + r"(?![A-Za-z0-9+#])", re.I)


def keyword_coverage(text, keywords):
    covered, missing = [], []
    seen = set()
    for kw in keywords or []:
        kw = str(kw).strip()
        if not kw or kw.lower() in seen or len(kw) > 60:
            continue
        seen.add(kw.lower())
        (covered if _kw_regex(kw).search(text or "") else missing).append(kw)
    return covered, missing


# ---------------------------------------------------------------- DOCX inspection
def _iter_text(el, qn):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def inspect_docx(path):
    """Collect the facts the checks need from a DOCX file."""
    import docx
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    doc = docx.Document(str(path))
    body = doc.element.body
    paras = [Paragraph(p, doc) for p in body.iter(qn("w:p"))]
    tb_paras = {id(p) for tb in body.iter(qn("w:txbxContent")) for p in tb.iter(qn("w:p"))}
    tbl_paras = {id(p) for t in body.iter(qn("w:tbl")) for p in t.iter(qn("w:p"))}
    lines = []
    for p in paras:
        txt = "".join(r.text for r in p.runs) or p.text
        if not txt.strip():
            continue
        style = (p.style.name if p.style is not None else "") or ""
        lines.append({"text": txt, "style": style, "in_table": id(p._p) in tbl_paras, "in_textbox": id(p._p) in tb_paras,
                      "bold": bool(p.runs) and all(r.bold for r in p.runs if r.text.strip())})
    fonts, sizes = set(), []
    for p in paras:
        for r in p.runs:
            if r.font.name:
                fonts.add(r.font.name)
            if r.font.size:
                sizes.append(r.font.size.pt)
    for st in doc.styles:
        try:
            if st.type == 1 and st.font is not None and st.font.name and st.name in {l["style"] for l in lines}:
                fonts.add(st.font.name)
        except Exception:
            pass
    hf_text = []
    for sec in doc.sections:
        for part in (sec.header, sec.footer, sec.first_page_header, sec.first_page_footer):
            try:
                hf_text.append(_iter_text(part._element, qn))
            except Exception:
                pass
    columns = 0
    for sec in doc.sections:
        cols = sec._sectPr.find(qn("w:cols")) if sec._sectPr is not None else None
        if cols is not None:
            try:
                columns = max(columns, int(cols.get(qn("w:num")) or 1))
            except ValueError:
                pass
    images = len(list(body.iter(qn("w:drawing")))) + len(list(body.iter(qn("w:pict"))))
    return {
        "lines": lines, "fonts": fonts, "min_size": min(sizes) if sizes else None,
        "tables": len(body.findall(qn("w:tbl"))), "nested_tables": sum(1 for t in body.iter(qn("w:tbl")) for _ in t.iter(qn("w:tbl"))) - len(body.findall(qn("w:tbl"))),
        "text_boxes": len(list(body.iter(qn("w:txbxContent")))), "images": images, "columns": columns,
        "header_footer_text": "\n".join(hf_text), "body_text": "\n".join(l["text"] for l in lines),
    }


def inspect_text(text):
    lines = [{"text": l.strip(), "style": "", "in_table": False, "in_textbox": False, "bold": False}
             for l in (text or "").splitlines() if l.strip()]
    return {"lines": lines, "fonts": set(), "min_size": None, "tables": 0, "nested_tables": 0, "text_boxes": 0,
            "images": 0, "columns": 0, "header_footer_text": "", "body_text": "\n".join(l["text"] for l in lines)}


# ---------------------------------------------------------------- checks
def _check(cid, label, status, detail):
    return {"id": cid, "label": label, "status": status, "detail": detail}


def run_checks(info, jd_keywords=(), fmt="docx"):
    """Return {score, structure_score, keyword_score, checks, covered, missing}."""
    body = info["body_text"]
    checks = []
    if fmt == "docx":
        checks.append(_check("filetype", "File type", "pass", "DOCX is read by every ATS."))
    elif fmt == "txt":
        checks.append(_check("filetype", "File type", "pass", "Plain text is read by every ATS."))

    if info["columns"] > 1:
        checks.append(_check("columns", "Single column layout", "fail", f"{info['columns']} columns found. Parsers read across columns and scramble the text."))
    else:
        checks.append(_check("columns", "Single column layout", "pass", "No multi-column sections."))

    if info["text_boxes"]:
        checks.append(_check("textboxes", "No text boxes", "fail", f"{info['text_boxes']} text box(es). Most parsers skip text inside them."))
    else:
        checks.append(_check("textboxes", "No text boxes", "pass", "No text boxes."))

    if info["nested_tables"] > 0:
        checks.append(_check("tables", "No tables", "fail", "Nested tables found. Content is usually lost or jumbled."))
    elif info["tables"]:
        checks.append(_check("tables", "No tables", "warn", f"{info['tables']} table(s). Simple tables often parse, but cells can be read out of order."))
    else:
        checks.append(_check("tables", "No tables", "pass", "No tables."))

    if info["images"]:
        checks.append(_check("images", "No images or graphics", "warn", f"{info['images']} image(s) or shape(s). ATS ignores them; icons used as labels lose their meaning."))
    else:
        checks.append(_check("images", "No images or graphics", "pass", "No images."))

    hf = info["header_footer_text"]
    hf_contact = bool(EMAIL.search(hf) or PHONE.search(hf))
    body_email, body_phone = bool(EMAIL.search(body)), bool(PHONE.search(body))
    if hf_contact and not (body_email and body_phone):
        checks.append(_check("contact", "Contact details in the body", "fail", "Email or phone sits in the page header or footer, which many parsers do not read."))
    elif not body_email:
        checks.append(_check("contact", "Contact details in the body", "fail", "No email address found in the resume body."))
    elif not body_phone:
        checks.append(_check("contact", "Contact details in the body", "warn", "No phone number found in the resume body."))
    else:
        checks.append(_check("contact", "Contact details in the body", "pass", "Email and phone are in the body text."))

    found = {s for l in info["lines"] if (s := section_of(l["text"]))}
    missing_secs = [SECTIONS[k][1] for k in REQUIRED_SECTIONS if k not in found]
    if missing_secs:
        checks.append(_check("headings", "Standard section headings", "warn", "Not recognised: " + ", ".join(missing_secs) + ". Parsers look for standard names such as Work Experience, Education and Skills."))
    else:
        checks.append(_check("headings", "Standard section headings", "pass", "Experience, Education and Skills headings recognised."))

    odd_fonts = sorted(f for f in info["fonts"] if f and f.lower() not in SAFE_FONTS and not f.lower().startswith(("symbol", "wingdings")))
    sym_fonts = sorted(f for f in info["fonts"] if f and f.lower().startswith(("symbol", "wingdings")))
    if sym_fonts:
        checks.append(_check("fonts", "Standard fonts", "warn", "Symbol fonts in use (" + ", ".join(sym_fonts) + "). Characters from them are read as junk."))
    elif odd_fonts:
        checks.append(_check("fonts", "Standard fonts", "warn", "Uncommon fonts: " + ", ".join(odd_fonts[:4]) + ". Safer: Calibri, Arial, Georgia or Times New Roman."))
    else:
        checks.append(_check("fonts", "Standard fonts", "pass", "Only common fonts are used." if info["fonts"] else "Default fonts."))

    if info["min_size"] is not None and info["min_size"] < 9:
        checks.append(_check("fontsize", "Readable font size", "warn", f"Smallest text is {info['min_size']:g} pt. Keep body text at 10 pt or larger."))
    else:
        checks.append(_check("fontsize", "Readable font size", "pass", "Text sizes are readable."))

    odd = sorted(set(ODD_CHARS.findall(body)))
    if odd:
        checks.append(_check("chars", "No special characters", "warn", "Unusual symbols found: " + " ".join(odd[:8]) + ". Replace icons and fancy symbols with plain text."))
    else:
        checks.append(_check("chars", "No special characters", "pass", "Only standard characters and bullets."))

    words = len(re.findall(r"\w+", body))
    if words > 1200:
        checks.append(_check("length", "Length", "warn", f"About {words} words, which is more than two pages. Trim older or unrelated detail."))
    elif words < 150:
        checks.append(_check("length", "Length", "warn", f"Only about {words} words. Parsers may not find enough to match."))
    else:
        checks.append(_check("length", "Length", "pass", f"About {words} words."))

    dated = sum(1 for l in info["lines"] if RANGE.search(l["text"]))
    if dated == 0:
        checks.append(_check("dates", "Employment dates", "warn", "No date ranges such as 'Jan 2021 - Present' were found. Parsers use them to compute experience."))
    else:
        checks.append(_check("dates", "Employment dates", "pass", f"{dated} date range(s) found."))

    covered, missing = keyword_coverage(body, jd_keywords)
    total = len(covered) + len(missing)
    kw_score = round(100 * len(covered) / total) if total else 100
    checks.append(_check("keywords", "Job keywords", "pass" if kw_score >= 75 else "warn" if kw_score >= 50 else "fail",
                         f"{len(covered)} of {total} key terms from the job appear in the resume." if total else "No job keywords were extracted."))

    fails = sum(1 for c in checks if c["status"] == "fail")
    warns = sum(1 for c in checks if c["status"] == "warn" and c["id"] != "keywords")
    structure = max(0, 100 - 20 * fails - 6 * warns)
    score = round(0.65 * structure + 0.35 * kw_score)
    return {"score": score, "structure_score": structure, "keyword_score": kw_score, "checks": checks,
            "covered": covered, "missing": missing}


def analyze(path_or_text, jd_keywords=(), fmt="docx"):
    info = inspect_docx(path_or_text) if fmt == "docx" else inspect_text(path_or_text)
    return run_checks(info, jd_keywords, fmt)


# ---------------------------------------------------------------- ATS-safe copy
def classify_lines(lines, header_footer_text=""):
    """Turn inspected lines into (kind, text) tuples: name, contact, heading, bullet, para."""
    out = []
    contact_from_hf = []
    for raw in (header_footer_text or "").splitlines():
        if EMAIL.search(raw) or PHONE.search(raw) or LINK.search(raw):
            contact_from_hf.append(raw.strip())
    name_done = False
    for l in lines:
        t = re.sub(r"\s{2,}", "  ", l["text"].strip())
        if not t:
            continue
        sec = section_of(t)
        if not name_done and len(t) <= 40 and re.fullmatch(r"[A-Za-z][A-Za-z .'\-]+", t) and not sec:
            out.append(("name", t))
            name_done = True
            continue
        if EMAIL.search(t) or LINK.search(t) or (PHONE.search(t) and len(t) < 120 and not RANGE.search(t)):
            out.append(("contact", t))
            continue
        if sec and sec != "other":
            out.append(("heading", SECTIONS[sec][1]))
            continue
        if sec == "other" or (l.get("style", "").lower().startswith("heading") and len(t) < 60):
            out.append(("heading", t.rstrip(":").title() if t.isupper() else t.rstrip(":")))
            continue
        m = BULLET.match(t)
        if m or l.get("style", "").lower().startswith("list"):
            out.append(("bullet", t[m.end():].strip() if m else t))
            continue
        out.append(("para", t))
    if contact_from_hf and not any(k == "contact" for k, _ in out):
        insert_at = 1 if out and out[0][0] == "name" else 0
        for c in reversed(contact_from_hf):
            out.insert(insert_at, ("contact", c))
    return out


def build_ats_docx(items, out_path):
    """Write a plain single-column DOCX: Calibri, standard headings, real bullets, no tables or boxes."""
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
    d = docx.Document()
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Pt(58)
        sec.top_margin = sec.bottom_margin = Pt(50)
    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.05
    for kind, text in items:
        text = ODD_CHARS.sub("", text).strip()
        if not text:
            continue
        if kind == "name":
            p = d.add_paragraph()
            r = p.add_run(text)
            r.bold = True
            r.font.size = Pt(18)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(1)
        elif kind == "contact":
            p = d.add_paragraph(text.replace(" | ", "  |  "))
            p.paragraph_format.space_after = Pt(2)
        elif kind == "heading":
            p = d.add_paragraph()
            r = p.add_run(text.upper())
            r.bold = True
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
        elif kind == "bullet":
            p = d.add_paragraph(text, style="List Bullet")
            p.paragraph_format.space_after = Pt(1)
        else:
            d.add_paragraph(text)
    d.save(str(out_path))
    return Path(out_path)
