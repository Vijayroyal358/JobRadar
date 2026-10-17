"""Resume parsing: text extraction, skills, total experience, current title."""
import io
import re
from datetime import datetime

from .skills import extract_skills, suggest_roles

TITLE_WORDS = re.compile(
    r"\b(engineer|developer|analyst|manager|administrator|admin|consultant|specialist|lead|architect|tester|"
    r"programmer|designer|scientist|technician|executive|associate|officer|intern|support|devops|sre|sdet|"
    r"head|director|coordinator|owner|master|expert|professional|trainee)\b", re.I)

EXP_HEAD = re.compile(r"^\s*(work\s+experience|professional\s+experience|experience|employment(\s+history)?|"
                      r"work\s+history|career\s+(history|summary)|professional\s+background|relevant\s+experience)\s*:?\s*$",
                      re.I)
OTHER_HEAD = re.compile(r"^\s*(education(al)?(\s+(background|qualifications?|details))?|academic(s|\s+\w+)?|qualifications?|"
                        r"certifications?|projects?|academic\s+projects|skills|technical\s+skills|key\s+skills|"
                        r"achievements|awards|languages|hobbies|interests|personal\s+(details|information)|declaration|"
                        r"trainings?|publications|references)\s*:?\s*$", re.I)

MON = (r"(jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|jul(?:y)?|aug(?:ust)?|"
       r"sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)")
_MON_IDX = {m: i + 1 for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}
DATE = rf"(?:{MON}[\s.,'’\-]*(\d{{4}}|\d{{2}})|(\d{{1,2}})[/\-.](\d{{4}})|((?:19|20)\d{{2}}))"
END = rf"(?:{DATE}|(present|current|now|till\s*date|to\s*date|date|ongoing|today))"
RANGE = re.compile(rf"{DATE}\s*(?:-|–|—|to|till|until|~)\s*{END}", re.I)
EXPLICIT_YEARS = re.compile(r"(\d{1,2}(?:\.\d{1,2})?)\s*\+?\s*(?:years?|yrs?)(?:\s+and\s+\d{1,2}\s+months?)?\s*(?:of\s+)?"
                            r"(?:[A-Za-z\-/&]+\s+){0,4}?(?:experience|exp\b)", re.I)
EDU_LINE = re.compile(r"\b(university|college|school|institute|b\.?\s?tech|m\.?\s?tech|bachelor|master|b\.?e\b|m\.?e\b|"
                      r"b\.?sc|m\.?sc|mba|mca|bca|diploma|ssc|hsc|cgpa|gpa|percentage|graduat)", re.I)


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        pages = []
        for p in reader.pages:
            try:
                pages.append(p.extract_text(extraction_mode="layout") or p.extract_text() or "")
            except Exception:
                pages.append(p.extract_text() or "")
        text = "\n".join(pages)
    elif name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    elif name.endswith((".txt", ".md")):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT resume.")
    text = text.replace(" ", " ").replace("–", "-").replace("—", "-")
    text = re.sub(r"[ \t]{2,}", "  ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _lines(text):
    return [l.strip() for l in text.splitlines()]


def _experience_section(text: str) -> str:
    lines = _lines(text)
    start = None
    for i, l in enumerate(lines):
        if EXP_HEAD.match(l):
            start = i + 1
            break
    if start is None:
        return ""
    out = []
    for l in lines[start:]:
        if OTHER_HEAD.match(l):
            break
        out.append(l)
    return "\n".join(out)


def _ym(mon, yr, num_mon, num_yr, year_only, is_end=False):
    if mon:
        y = int(yr)
        y = y + 2000 if y < 100 and y <= datetime.now().year % 100 + 1 else (y + 1900 if y < 100 else y)
        return y * 12 + _MON_IDX[mon[:3].lower()] - 1
    if num_mon and num_yr:
        m = int(num_mon)
        if 1 <= m <= 12:
            return int(num_yr) * 12 + m - 1
        return None
    if year_only:
        return int(year_only) * 12 + (11 if is_end else 0)
    return None


def _months_from_ranges(text: str) -> int:
    now = datetime.now()
    now_m = now.year * 12 + now.month - 1
    intervals = []
    for line in text.splitlines():
        if EDU_LINE.search(line):
            continue
        for m in RANGE.finditer(line):
            g = m.groups()
            start = _ym(g[0], g[1], g[2], g[3], g[4])
            if g[10]:
                end = now_m
            else:
                end = _ym(g[5], g[6], g[7], g[8], g[9], is_end=True)
            if start is None or end is None:
                continue
            if not (1970 * 12 <= start <= now_m) or end < start:
                continue
            end = min(end, now_m)
            intervals.append((start, end))
    if not intervals:
        return 0
    intervals.sort()
    merged = [list(intervals[0])]
    for s, e in intervals[1:]:
        if s <= merged[-1][1] + 1:
            merged[-1][1] = max(merged[-1][1], e)
        else:
            merged.append([s, e])
    return sum(e - s + 1 for s, e in merged)


def estimate_years(text: str):
    """Returns (years, method). Prefers an explicit 'N years of experience' statement."""
    head = text[:2500]
    explicit = [float(m.group(1)) for m in EXPLICIT_YEARS.finditer(head) if float(m.group(1)) <= 40]
    section = _experience_section(text)
    months = _months_from_ranges(section) if section else 0
    if not months:
        months = _months_from_ranges(text)
    from_ranges = round(months / 12, 1) if months else None
    if explicit:
        return max(explicit), "stated in resume"
    if from_ranges:
        return from_ranges, "calculated from job dates"
    return None, "not found"


def guess_title(text: str) -> str:
    m = re.search(r"(?:current\s+)?(?:designation|position|role|title)\s*[:\-]\s*([^\n|]{3,70})", text, re.I)
    if m and TITLE_WORDS.search(m.group(1)):
        return m.group(1).strip(" ,.-")
    section = _experience_section(text)
    for l in _lines(section)[:15]:
        for chunk in re.split(r"\s{2,}|\||,|\s-\s|\bat\b", l):
            c = chunk.strip(" ,.-:")
            if 3 <= len(c) <= 60 and TITLE_WORDS.search(c) and not RANGE.search(c) and not EDU_LINE.search(c):
                return c
    for l in _lines(text)[1:10]:
        if 3 <= len(l) <= 70 and TITLE_WORDS.search(l) and "@" not in l and not EDU_LINE.search(l):
            return re.split(r"\s{2,}|\|", l)[0].strip(" ,.-")
    return ""


def guess_name(text: str) -> str:
    for l in _lines(text)[:6]:
        l = re.split(r"\s{2,}|\|", l)[0].strip()
        if 3 <= len(l) <= 40 and re.fullmatch(r"[A-Za-z][A-Za-z .'\-]+", l) and not TITLE_WORDS.search(l) \
                and l.lower() not in ("resume", "curriculum vitae", "cv"):
            return l.title() if l.isupper() else l
    return ""


def parse_resume(filename: str, data: bytes) -> dict:
    text = extract_text(filename, data)
    if len(text) < 80:
        raise ValueError("Could not read text from this file. If it is a scanned image PDF, export it as a text PDF or DOCX.")
    skills = sorted(extract_skills(text))
    years, method = estimate_years(text)
    title = guess_title(text)
    return {
        "name": guess_name(text),
        "current_title": title,
        "years_experience": years,
        "years_method": method,
        "skills": skills,
        "suggested_roles": suggest_roles(skills),
        "text": text,
    }
